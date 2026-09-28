import io

from openpyxl import Workbook, load_workbook


def test_export_excel_contains_all_cards_with_current_state(client):
    client.post("/api/inventory/UNL-003/increment", params={"delta": 2})
    client.post("/api/inventory/UNL-003/increment", params={"foil": True})
    client.patch("/api/inventory/UNL-006", json={"limit_override": 5, "in_binder": True})

    resp = client.get("/api/backup/excel")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    ws = wb.active
    assert ws.title == "Collection"
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[:5] == ["Card ID", "Set", "Number", "Name", "Rarity"]
    assert header[5:] == [
        "Count", "Foil count", "In binder", "Binder foil", "Limit override",
    ]

    by_id = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert by_id["UNL-003"][5] == 2 and by_id["UNL-003"][6] == 1  # count / foil
    assert by_id["UNL-006"][7] is True and by_id["UNL-006"][9] == 5  # in_binder / override
    # Every seed card appears, prefilled with zeros when there's no inventory.
    all_cards = client.get("/api/cards").json()
    assert len(by_id) == len(all_cards)
    zero_row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "UNL-005")
    assert zero_row[5:9] == (0, 0, False, False)


def test_import_excel_applies_edits_and_skips_unknown(client):
    resp = client.get("/api/backup/excel")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    # Find the UNL-004 row and edit it as if the user did in Excel.
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "UNL-004":
            row[5].value = 3  # count
            row[6].value = 1  # foil count
            row[7].value = True  # in binder
            row[9].value = 7  # limit override
            break
    # Add an unknown card row (should be counted as skipped).
    ws.append(["ZZZ-999", "ZZZ", "1", "Not a real card", "Common", 4, 0, False, False, None])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    up = client.post(
        "/api/backup/excel",
        files={
            "file": (
                "edited.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert up.status_code == 200
    result = up.json()
    assert result["skipped"] == 1  # the ZZZ row
    assert result["applied"] >= 1

    card = {c["id"]: c for c in client.get("/api/cards").json()}["UNL-004"]
    assert card["count"] == 3 and card["foil_count"] == 1
    assert card["in_binder"] is True
    assert card["limit"] == 7  # limit_override took effect


def test_import_excel_enforces_binder_foil_implies_in_binder(client):
    # Build a minimal sheet: user marked binder_foil=True but forgot in_binder.
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Card ID", "Set", "Number", "Name", "Rarity",
        "Count", "Foil count", "In binder", "Binder foil", "Limit override",
    ])
    ws.append(["UNL-007", "UNL", "007", "LeBlanc", "Epic", 0, 0, False, True, None])
    buf = io.BytesIO()
    wb.save(buf)
    resp = client.post(
        "/api/backup/excel",
        files={"file": ("t.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    card = {c["id"]: c for c in client.get("/api/cards").json()}["UNL-007"]
    assert card["binder_foil"] is True and card["in_binder"] is True


def test_import_excel_accepts_string_boolean_and_blank_override(client):
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Card ID", "Set", "Number", "Name", "Rarity",
        "Count", "Foil count", "In binder", "Binder foil", "Limit override",
    ])
    # Blank limit override, "TRUE" string, numeric string count.
    ws.append(["UNL-008", "UNL", "008", "Calculated Approach", "Uncommon", "4", 0, "TRUE", "no", ""])
    buf = io.BytesIO()
    wb.save(buf)
    resp = client.post(
        "/api/backup/excel",
        files={"file": ("t.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    card = {c["id"]: c for c in client.get("/api/cards").json()}["UNL-008"]
    assert card["count"] == 4
    assert card["in_binder"] is True
    assert card["binder_foil"] is False
    assert card["limit_override"] is None


def test_import_excel_rejects_non_xlsx_and_garbage(client):
    r = client.post("/api/backup/excel",
                    files={"file": ("notes.txt", b"hello", "text/plain")})
    assert r.status_code == 415
    r = client.post("/api/backup/excel",
                    files={"file": ("broken.xlsx", b"not really xlsx",
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 422
